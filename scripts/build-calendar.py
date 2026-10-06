#!/usr/bin/env python3
"""
Regenerate content/calendar-data.ts from the Tallahassee youth arts spreadsheet.

    python3 scripts/build-calendar.py path/to/Tallahassee_Youth_Arts_Events_2026-27.xlsx

Needs openpyxl (pip3 install openpyxl).

This repo is public, so the spreadsheet itself is NOT committed, and these
are deliberately never carried over: the Contact column, the Contacts sheet,
the Notes columns (internal research), and the Ongoing Programs sheet.
Links to mrgermans.weebly.com (a teacher's personal calendar page, cited as
the only source for many events) are dropped rather than published.
"""

import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "content" / "calendar-data.ts"
DROPPED_LINK_HOSTS = ("mrgermans.weebly.com",)


def iso(value):
    if value is None or value == "":
        return None
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def disciplines(raw):
    s = (raw or "").lower()
    out = []
    if "band" in s:
        out.append("Band")
    if "choral" in s or "choir" in s:
        out.append("Choral")
    if "orchestra" in s:
        out.append("Orchestra")
    if "theatre" in s:
        out.append("Theatre")
    if "dance" in s:
        out.append("Dance")
    if "visual" in s:
        out.append("Visual Art")
    if "music" in s:
        out.append("Other Music")
    if "multi" in s or not out:
        out.append("Multi-arts")
    return out


def clean_time(t):
    if not t:
        return None
    t = str(t).strip()
    if t.upper().startswith("TBA"):  # includes research caveats like "TBA (past concerts 3 PM)"
        return None
    if t == "Daily (hours TBA)":
        return "Daily"
    return re.sub(r"(?<=[\dM])-(?=\d)", "–", t)


def clean_venue(v):
    if not v:
        return None
    v = str(v).strip()
    if v.upper() == "TBA" or v.lower().startswith("location not listed"):
        return None
    m = re.match(r"^TBA \((.+)\)$", v)
    if m:
        return "FSU campus (venue TBA)" if m.group(1) == "FSU campus" else None
    return v


def clean_cost(c):
    """Returns (cost_text_or_None, is_free)."""
    if not c:
        return None, False
    c = str(c).strip()
    if c.lower() == "free":
        return None, True
    if c.upper().startswith("TBA") or "past" in c.lower():
        return None, False
    return c, False


def clean_url(u):
    if not u:
        return None
    u = str(u).strip()
    if not u.startswith("http") or any(h in u for h in DROPPED_LINK_HOSTS):
        return None
    return u


def undated_timing(time_text, title):
    t = (time_text or "").strip()
    if t.upper().startswith("TBA"):
        m = re.search(r"\((early [A-Za-z]+)\)", title)
        return m.group(1).capitalize() if m else "Date to be announced"
    return re.split(r"\s*\(", t)[0] or "Date to be announced"


def strip_none(d):
    return {k: v for k, v in d.items() if v not in (None, [], "")}


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    dated, undated = [], []

    # Events sheet: # | Start | End | Days | Time | Event | Host | Discipline | Participants |
    #               Type | Venue | Cost | Contact | Date Status | Source | Notes
    for r in list(wb["Events"].iter_rows(values_only=True))[1:]:
        if not r[5]:
            continue
        cost, free = clean_cost(r[11])
        base = dict(
            title=str(r[5]).strip(),
            host=str(r[6]).strip() if r[6] else "",
            disciplines=disciplines(r[7]),
            venue=clean_venue(r[10]),
            cost=cost,
            free=True if free else None,
            url=clean_url(r[14]),
        )
        start, end = iso(r[1]), iso(r[2])
        if not start:
            undated.append(
                strip_none(dict(id=f"y{r[0]}", timing=undated_timing(r[4], base["title"]), **base))
            )
            continue
        dated.append(
            strip_none(
                dict(
                    id=f"y{r[0]}",
                    start=start,
                    end=end if end and end != start else None,
                    time=clean_time(r[4]),
                    confirmed=str(r[13] or "").startswith("Confirmed"),
                    **base,
                )
            )
        )

    # Regional & State sheet: Start | End | Event | Org | Discipline | Location | Notes | Source
    for i, r in enumerate(list(wb["Regional & State"].iter_rows(values_only=True))[1:], 1):
        if not r[2]:
            continue
        start, end = iso(r[0]), iso(r[1])
        dated.append(
            strip_none(
                dict(
                    id=f"r{i}",
                    title=str(r[2]).strip(),
                    start=start,
                    end=end if end and end != start else None,
                    host=str(r[3]).strip() if r[3] else "",
                    disciplines=disciplines(r[4]),
                    venue=clean_venue(r[5]),
                    confirmed=False,
                    regional=True,
                    url=clean_url(r[7]),
                )
            )
        )

    dated.sort(key=lambda e: (e["start"], e["id"]))

    today = datetime.date.today().isoformat()
    lines = [
        "// GENERATED by scripts/build-calendar.py — do not edit by hand.",
        "// Source: Tallahassee_Youth_Arts_Events_2026-27.xlsx (Events + Regional & State sheets).",
        "// Contacts, internal notes, and the Ongoing Programs sheet are intentionally left out.",
        'import type { CalendarEvent, UndatedEvent } from "@/lib/calendar";',
        "",
        f'export const calendarUpdated = "{today}";',
        "",
        "export const calendarEvents: CalendarEvent[] = [",
        *[f"  {json.dumps(e, ensure_ascii=False)}," for e in dated],
        "];",
        "",
        "export const undatedEvents: UndatedEvent[] = [",
        *[f"  {json.dumps(e, ensure_ascii=False)}," for e in undated],
        "];",
        "",
    ]
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {OUT}: {len(dated)} dated events, {len(undated)} undated.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
